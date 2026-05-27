import base64
import json
import math
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover - handled at runtime
    load_workbook = None

try:
    import xlrd
except Exception:  # pragma: no cover - handled at runtime
    xlrd = None


DATA_DIR = Path(__file__).resolve().parent / "local_data"
STATE_PATH = DATA_DIR / "apps_script_state.json"

TIPOS_PRODUCTO = ["Cabeza", "Patas y Manos", "Visceras Blancas", "Visceras Rojas"]
TURNOS = ["SxD", "VxS", "JxV", "MxJ", "MxM", "LxM", "DxL"]
OPL_DEFAULT = "TRANSCARNES"
OPL_EXCEPCIONES_DEFAULT = [
    ["AVILA MONSALVE REINALDO", "DRA CAVA", 0],
    ["BENITEZ GARNICA CEFERINO", "EDGAR AM", 0],
    ["CALIXTO ARDILA JAIME", "DRA CAVA", 0],
    ["CARNES SANTACRUZ S.A.S", "CSZ B/GA", 0],
    ["CRUZ LEONIDAS", "CAVA WO", 0],
    ["DRISTRIBUDORA DE CARNES AJR S.A.S", "CAVA AJR", 0],
    ["INVERSIONES ZULUAGA RUEDA S.A.S.", "MLT. GUARIN", 0],
    ["JAIMES BERMUDEZ JOSE MARIA", "MLT. GUARIN", 0],
    ["SANCHEZ CALDERON MIREYA", "CAVA MIREYA", 0],
    ["SUPERMERCADOS MAS POR MENOS S.A.S.", "MLT. GUARIN", 0],
    ["TECNOLOGIAS AGROPECUARIAS DE COLOMBIA S.A.S.", "CAVA T.A", 0],
    ["ROMERO OSORIO JOHN IGNACIO", "SMOYA", 0],
    ["COLBEEF S.A.S", "MLT. GUARIN", 0],
]
PUESTOS_EXCLUIDOS_DESP = {
    "01305/TEMP1 /DxL///CALLE 23# 6-52 PLACITA GIRARDOT",
    "03105/Guarin //Cra 33a # 32-109",
    "05200/TEMP1 /DxL///CARRERA 3#61-39 LOS NARANJOS",
    "12157/Giron /DxL///CARRERA 22 # 13a-10 barrio EL CONSUELO",
    "379P/Piedecuesta /VxS/",
    "CAVA AJR/CAVA //CAVA AJR",
    "CAVA FORTUNATO/CAVA //CAVA",
    "CAVA MIREYA/CAVA //CAVA",
    "CAVA./CAVA ///",
    "CCARNES CAVA/CARNES Y CARNES //Cra 34 W #71-100 Bdga 46",
    "OLIMPICA/Barranquilla //6ta Entrada Km 2-701 via Caracoli-Malambo",
    "OLIMPICA/Barranquilla //6ta Entrada Km 2-701 vía Caracolí-Malambo",
    "RH32/Temp 2 Giron /DxL///CRA 29 # 33-70 LLANITO PARTE BAJA",
    "RH32/Temp 2 Girón /DxL///CRA 29 # 33-70 LLANITO PARTE BAJA",
}
CAVAS_DEFAULT = [
    {"grupo": "V. Rojas & Blancas (V.Rojas)", "carros": 40, "capPorCarro": 20, "inventario": 0},
    {"grupo": "V. Rojas & Blancas (V.Blancas)", "carros": 22, "capPorCarro": 25, "inventario": 0},
    {"grupo": "V. Acondicionamiento", "carros": 22, "capPorCarro": 25, "inventario": 0},
    {"grupo": "Patas & Manos", "carros": 80, "capPorCarro": 9, "inventario": 0},
    {"grupo": "Cabezas", "carros": 80, "capPorCarro": 9, "inventario": 0},
]
PERCHEROS_DEFAULT = [
    {"cava": "V. Rojas & Blancas", "blancas": 0, "rojas": 0, "patasManos": 0, "cabezas": 0, "crudas": 0},
    {"cava": "V. Acondicionamiento", "blancas": 0, "rojas": 0, "patasManos": 0, "cabezas": 0, "crudas": 0},
    {"cava": "Patas & Cabezas", "blancas": 0, "rojas": 0, "patasManos": 0, "cabezas": 0, "crudas": 0},
    {"cava": "Recepción", "blancas": 0, "rojas": 0, "patasManos": 0, "cabezas": 0, "crudas": 0},
    {"cava": "Retenidos", "blancas": 0, "rojas": 0, "patasManos": 0, "cabezas": 0, "crudas": 0},
]


def _today():
    return datetime.now().strftime("%d/%m/%Y")


def _now():
    return datetime.now().strftime("%d/%m/%Y %H:%M")


def _blank_state():
    return {
        "sheets": {},
        "resumen_decomisos": [],
        "resumen_despachos": {"hayDatos": False},
        "historial_pdf": [],
        "informe": {
            "fecha": _today(),
            "completos": 0,
            "incompletos": 0,
            "beneficioDia": 0,
            "stockTotal": 100,
            "danados": 2,
            "cavas": CAVAS_DEFAULT,
            "novedades": [],
            "percheros": PERCHEROS_DEFAULT,
        },
        "opl_config": OPL_EXCEPCIONES_DEFAULT,
        "opl_progreso": [],
        "historico": [],
        "planilla": {},
    }


def _load_state():
    DATA_DIR.mkdir(exist_ok=True)
    if not STATE_PATH.exists():
        state = _blank_state()
        _save_state(state)
        return state
    with STATE_PATH.open("r", encoding="utf-8") as fh:
        state = json.load(fh)
    base = _blank_state()
    for key, value in base.items():
        state.setdefault(key, value)
    return state


def _save_state(state):
    DATA_DIR.mkdir(exist_ok=True)
    with STATE_PATH.open("w", encoding="utf-8") as fh:
        json.dump(state, fh, ensure_ascii=False, indent=2, default=str)


def _as_str(value):
    return "" if value is None else str(value).strip()


def _num(value):
    try:
        if value in ("", None):
            return 0
        return float(value)
    except Exception:
        return 0


def _codigo_base(value):
    s = re.sub(r"[^0-9\-]", "", _as_str(value))
    pos = s.rfind("-")
    return s[:pos] if pos > 0 else s


def _detectar_turno_por_dia():
    # Python weekday: Monday=0. Apps Script mapping: Sunday=DxL, Monday=LxM...
    return {0: "LxM", 1: "MxM", 2: "MxJ", 3: "JxV", 4: "VxS", 5: "SxD", 6: "DxL"}[datetime.now().weekday()]


def _rows_from_range(rows, start_row_1, start_col_1, width):
    start_r = max(0, start_row_1 - 1)
    start_c = max(0, start_col_1 - 1)
    out = []
    for row in rows[start_r:]:
        padded = list(row) + [""] * (start_c + width - len(row))
        out.append(padded[start_c:start_c + width])
    return out


def _read_workbook_from_base64(payload):
    raw = base64.b64decode(payload or "")
    return _read_excel_bytes(raw)


def _read_workbook_from_bytes(byte_values):
    raw = bytes(byte_values or [])
    return _read_excel_bytes(raw)


def _read_excel_bytes(raw):
    if load_workbook is not None:
        try:
            return _read_xlsx_bytes(raw)
        except Exception:
            pass
    if xlrd is not None:
        try:
            return _read_xls_bytes(raw)
        except Exception:
            pass
    raise RuntimeError("No se pudo leer el Excel. Usa .xlsx o instala openpyxl/xlrd para .xls.")


def _read_xlsx_bytes(raw):
    wb = load_workbook(BytesIO(raw), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if cell is None else cell for cell in row])
    return rows


def _read_xls_bytes(raw):
    book = xlrd.open_workbook(file_contents=raw)
    sheet = book.sheet_by_index(0)
    return [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]


def _opl_map(state):
    config = state.get("opl_config") or []
    return {_as_str(row[0]).upper(): (_as_str(row[1]) or OPL_DEFAULT) for row in config if row and _as_str(row[0])}


def _resolver_opl(prop, mapa):
    return mapa.get(_as_str(prop).upper(), OPL_DEFAULT)


def initializeSheets():
    state = _load_state()
    _save_state(state)
    return {"success": True}


def importarExcel(base64_payload, sheetName):
    try:
        state = _load_state()
        rows = _read_workbook_from_base64(base64_payload)
        state["sheets"][sheetName] = rows
        _save_state(state)
        return {"success": True, "rows": len(rows)}
    except Exception as exc:
        return {"success": False, "message": str(exc)}


def importarExcelAdicionales(byte_values, nombreArchivo):
    try:
        state = _load_state()
        rows = _read_workbook_from_bytes(byte_values)
        state["sheets"]["Adicionales_Temp"] = rows
        state["adicionales_nombre"] = nombreArchivo
        _save_state(state)
        return {"success": True, "rows": len(rows)}
    except Exception as exc:
        return {"success": False, "message": str(exc)}


def importarAdicionales(_fileData=None, nombreArchivo=None, _tipoManual=None):
    return {"success": True, "totalAdicional": 0, "totalCancel": 0, "totalCambio": 0, "nombreArchivo": nombreArchivo or ""}


def contarJuegosViscerales():
    state = _load_state()
    rows = state["sheets"].get("Estado_Cavas", [])
    valores = _rows_from_range(rows, 13, 2, 1)
    bases = {_codigo_base(row[0]) for row in valores if _codigo_base(row[0])}
    return {"success": True, "total": len(bases)}


def _es_cruda(value):
    return _as_str(value).upper().startswith("CRUDAS")


def contarCrudas():
    state = _load_state()
    rows = state["sheets"].get("Estado_Cavas", [])
    datos = _rows_from_range(rows, 12, 2, 14)
    codigos = {}
    for row in datos:
        codigo, desc, col_o = _as_str(row[0]), _as_str(row[1]), row[13] if len(row) > 13 else ""
        if codigo and desc == "Visceras Blancas" and _es_cruda(col_o):
            base = _codigo_base(codigo)
            if base:
                codigos[base] = True
    return {"success": True, "total": len(codigos), "codigos": list(codigos)}


def resumirDecomisos():
    try:
        state = _load_state()
        data_cavas = _rows_from_range(state["sheets"].get("Estado_Cavas", []), 12, 2, 14)
        decomisos = state["sheets"].get("Reporte_Decomisos", [])
        destinos_por_base = {}
        for row in data_cavas:
            codigo = _as_str(row[0])
            base = _codigo_base(codigo)
            if base and base not in destinos_por_base:
                destinos_por_base[base] = _as_str(row[8] if len(row) > 8 else "")

        decomisos_por_base = {}
        for idx, row in enumerate(decomisos):
            if idx == 0:
                continue
            codigo = _as_str(row[0] if len(row) > 0 else "")
            producto = _as_str(row[2] if len(row) > 2 else "")
            base = _codigo_base(codigo)
            if base and producto:
                decomisos_por_base.setdefault(base, []).append({
                    "codigo": codigo,
                    "producto": producto,
                })

        resultados = []
        vistos = set()
        for base, items in decomisos_por_base.items():
            if base not in destinos_por_base:
                continue
            for item in items:
                key = (item["codigo"], item["producto"])
                if key in vistos:
                    continue
                vistos.add(key)
                resultados.append({
                    "id": item["codigo"],
                    "destino": destinos_por_base.get(base, ""),
                    "producto": item["producto"],
                })

        state["resumen_decomisos"] = resultados
        _actualizar_cantidades_iniciales_opl(state, data_cavas)
        _save_state(state)
        return {
            "success": True,
            "totalProductos": len(resultados),
            "totalDestinos": len({r["destino"] for r in resultados if r["destino"]}),
            "fechaProcesamiento": _now(),
            "resultados": resultados,
        }
    except Exception as exc:
        return {"success": False, "message": str(exc)}


def getResumenDecomisos():
    state = _load_state()
    resultados = state.get("resumen_decomisos", [])
    return {
        "success": True,
        "totalProductos": len(resultados),
        "totalDestinos": len({r.get("destino") for r in resultados if r.get("destino")}),
        "fechaProcesamiento": _now() if resultados else "Sin datos",
        "resultados": resultados,
    }


def limpiarResumen():
    state = _load_state()
    state["resumen_decomisos"] = []
    _save_state(state)
    return {"success": True}


def _detectar_turno_desde_datos(state):
    rows = _rows_from_range(state["sheets"].get("Despachos_Cavas", []), 15, 10, 1)
    counts = {t: 0 for t in TURNOS}
    for row in rows[:300]:
        puesto = _as_str(row[0])
        for turno in TURNOS:
            if turno in puesto:
                counts[turno] += 1
    best = max(counts.items(), key=lambda item: item[1])
    return best[0] if best[1] > 0 else _detectar_turno_por_dia()


def getTurnoActual():
    return {"success": True, "turno": _detectar_turno_desde_datos(_load_state())}


def procesarDespachos(turnoForzado=None):
    try:
        state = _load_state()
        raw_rows = state["sheets"].get("Despachos_Cavas", [])
        if len(raw_rows) < 15:
            return {"success": False, "message": "Hoja Despachos_Cavas no encontrada o sin datos desde fila 15."}
        rows = _rows_from_range(raw_rows, 15, 1, 13)
        turno = turnoForzado or _detectar_turno_desde_datos(state)

        def build_summary(use_turno=True, use_exclusions=True):
            mapa = {}
            for row in rows:
                codigo = _as_str(row[3] if len(row) > 3 else "")
                tipo = _as_str(row[7] if len(row) > 7 else "")
                puesto = _as_str(row[9] if len(row) > 9 else "")
                if not codigo or not tipo or not puesto:
                    continue
                # Algunos archivos solo dicen "CAVA ..." y no traen marca DxL/LxM/etc.
                # En ese caso no se debe descartar todo por turno.
                if use_turno and turno and any(t in puesto for t in TURNOS) and turno not in puesto:
                    continue
                if use_exclusions and puesto in PUESTOS_EXCLUIDOS_DESP:
                    continue
                mapa.setdefault(puesto, {tipo: 0 for tipo in TIPOS_PRODUCTO})
                if tipo in TIPOS_PRODUCTO:
                    mapa[puesto][tipo] += 1
            return mapa

        mapa_puestos = build_summary(use_turno=True, use_exclusions=True)
        if not mapa_puestos:
            mapa_puestos = build_summary(use_turno=False, use_exclusions=False)

        resultado = []
        for puesto in sorted(mapa_puestos):
            item = {"puesto": puesto}
            item.update(mapa_puestos[puesto])
            resultado.append(item)
        total_juegos = sum(item.get("Visceras Rojas", 0) for item in resultado)
        resumen = {
            "success": True,
            "hayDatos": bool(resultado),
            "turno": turno,
            "fechaUltima": _now(),
            "totalJuegos": total_juegos,
            "totalPuestos": len(resultado),
            "tipos": TIPOS_PRODUCTO,
            "resultado": resultado,
        }
        state["resumen_despachos"] = resumen
        _calcular_progreso_opl(state, total_juegos)
        _save_state(state)
        return resumen
    except Exception as exc:
        return {"success": False, "message": str(exc)}


def getResumenDespachoActual():
    state = _load_state()
    return state.get("resumen_despachos") or {"success": True, "hayDatos": False}


def getDetallesPuesto(puesto):
    state = _load_state()
    rows = _rows_from_range(state["sheets"].get("Despachos_Cavas", []), 15, 1, 13)
    filas = []
    for row in rows:
        if _as_str(row[9] if len(row) > 9 else "") == _as_str(puesto):
            filas.append({
                "id": _as_str(row[3] if len(row) > 3 else ""),
                "propietario": _as_str(row[4] if len(row) > 4 else ""),
                "tipo": _as_str(row[7] if len(row) > 7 else ""),
            })
    return {"success": True, "filas": filas}


def getDashboardDataDespachos():
    resumen = getResumenDespachoActual()
    return {
        "success": True,
        "totalJuegosDespachar": resumen.get("totalJuegos", 0) if resumen.get("hayDatos") else 0,
        "turnoDespacho": resumen.get("turno", ""),
        "ultimaActDespachos": resumen.get("fechaUltima", ""),
    }


def getDashboardData():
    try:
        resumen_decomisos = getResumenDecomisos()
        salidas = contarJuegosViscerales().get("total", 0)
        crudas = contarCrudas().get("total", 0)
        desp = getDashboardDataDespachos()
        pendientes = desp.get("totalJuegosDespachar", 0) or 0
        despachados = max(0, salidas - pendientes)
        progreso = min(100, round((despachados / salidas) * 100)) if salidas else 0
        return {
            "success": True,
            "totalSalidas": salidas,
            "totalDecomisos": resumen_decomisos.get("totalProductos", 0),
            "totalCrudas": crudas,
            "totalJuegosDespachar": pendientes,
            "turnoDespacho": desp.get("turnoDespacho", ""),
            "ultimaActDespachos": desp.get("ultimaActDespachos", ""),
            "progreso": progreso,
            "meta": salidas,
        }
    except Exception as exc:
        return {"success": False, "message": str(exc)}


def limpiarDespachos():
    state = _load_state()
    state["resumen_despachos"] = {"success": True, "hayDatos": False}
    state["opl_progreso"] = []
    _save_state(state)
    return {"success": True}


def getPuestosCrudas():
    state = _load_state()
    crudas = contarCrudas()
    set_crudas = set(crudas.get("codigos", []))
    resumen = state.get("resumen_despachos") or {}
    turno = resumen.get("turno", "")
    rows = _rows_from_range(state["sheets"].get("Despachos_Cavas", []), 15, 1, 13)
    puestos = {}
    for row in rows:
        codigo = _as_str(row[3] if len(row) > 3 else "")
        tipo = _as_str(row[7] if len(row) > 7 else "")
        puesto = _as_str(row[9] if len(row) > 9 else "")
        if turno and turno not in puesto:
            continue
        if tipo == "Visceras Blancas" and _codigo_base(codigo) in set_crudas:
            puestos[puesto] = True
    return {"success": True, "puestos": puestos, "total": crudas.get("total", 0), "codigos": crudas.get("codigos", [])}


def getCrudasDetalle():
    state = _load_state()
    rows = _rows_from_range(state["sheets"].get("Estado_Cavas", []), 12, 2, 14)
    mapa = _opl_map(state)
    crudas = {}
    for row in rows:
        codigo = _as_str(row[0])
        desc = _as_str(row[1])
        cliente = _as_str(row[3] if len(row) > 3 else "")
        puesto = _as_str(row[8] if len(row) > 8 else "")
        col_o = row[13] if len(row) > 13 else ""
        if not codigo or desc != "Visceras Blancas" or not _es_cruda(col_o):
            continue
        base = _codigo_base(codigo)
        key = f"{base}||{puesto}"
        if key not in crudas:
            crudas[key] = {
                "codigo": codigo,
                "base": base,
                "puesto": puesto,
                "cliente": cliente,
                "opl": _resolver_opl(cliente, mapa),
                "cantidad": 0,
                "observacion": _as_str(col_o).split("\n")[0],
            }
        crudas[key]["cantidad"] += 1
    return {"success": True, "filas": sorted(crudas.values(), key=lambda r: r.get("puesto", ""))}


def getCrudasResumenPDF():
    detalle = getCrudasDetalle()
    grupos = {}
    for row in detalle.get("filas", []):
        key = f"{row.get('puesto')}||{row.get('opl')}"
        grupos.setdefault(key, {"puesto": row.get("puesto", ""), "opl": row.get("opl", ""), "cantidad": 0, "codigos": []})
        grupos[key]["cantidad"] += row.get("cantidad", 0)
        grupos[key]["codigos"].append(row.get("codigo", ""))
    return {"success": True, "filas": sorted(grupos.values(), key=lambda r: (r["puesto"], r["opl"]))}


def _actualizar_cantidades_iniciales_opl(state, data_cavas):
    config = state.get("opl_config") or OPL_EXCEPCIONES_DEFAULT.copy()
    config_map = {_as_str(row[0]).upper(): row for row in config if row}
    conteo = {}
    vistos = set()
    for row in data_cavas:
        codigo = _as_str(row[0])
        tipo = _as_str(row[1] if len(row) > 1 else "")
        prop = _as_str(row[3] if len(row) > 3 else "")
        if not codigo or not prop or tipo != "Visceras Rojas":
            continue
        base = _codigo_base(codigo)
        if not base or base in vistos:
            continue
        vistos.add(base)
        conteo[prop.upper()] = conteo.get(prop.upper(), 0) + 1

    for row in config:
        while len(row) < 3:
            row.append(0)
        row[2] = 0
    for prop, total in conteo.items():
        if prop in config_map:
            config_map[prop][2] = total
        else:
            config.append([prop, OPL_DEFAULT, total])
    state["opl_config"] = config


def _calcular_progreso_opl(state, total_juegos=None):
    mapa = _opl_map(state)
    config = state.get("opl_config") or []
    totals_by_opl = {}
    for row in config:
        if len(row) < 3:
            continue
        opl = _as_str(row[1]) or OPL_DEFAULT
        total = int(_num(row[2]))
        if total > 0:
            totals_by_opl[opl] = totals_by_opl.get(opl, 0) + total

    if total_juegos is not None and int(total_juegos) == 0:
        state["opl_progreso"] = []
        state["operacion_finalizada"] = True
        return

    resumen = state.get("resumen_despachos") or {}
    turno = resumen.get("turno") or ""
    rows = _rows_from_range(state["sheets"].get("Despachos_Cavas", []), 15, 1, 13)
    pendientes_por_opl = {}
    vistos = set()
    for row in rows:
        codigo = _as_str(row[3] if len(row) > 3 else "")
        prop = _as_str(row[4] if len(row) > 4 else "")
        tipo = _as_str(row[7] if len(row) > 7 else "")
        puesto = _as_str(row[9] if len(row) > 9 else "")
        if tipo != "Visceras Rojas" or not codigo or not prop:
            continue
        if turno and turno not in puesto:
            continue
        base = _codigo_base(codigo)
        if not base or base in vistos:
            continue
        vistos.add(base)
        opl = _resolver_opl(prop, mapa)
        pendientes_por_opl[opl] = pendientes_por_opl.get(opl, 0) + 1

    progreso = []
    for opl, total in totals_by_opl.items():
        pendientes = min(total, pendientes_por_opl.get(opl, 0))
        despachados = max(0, total - pendientes)
        pct = round((despachados / total) * 100) if total else 100
        if total > 0 and pct < 100:
            progreso.append({
                "opl": opl,
                "total": total,
                "despachados": despachados,
                "pendientes": pendientes,
                "progreso": pct,
                "fecha": _now(),
            })
    state["opl_progreso"] = sorted(progreso, key=lambda r: r["pendientes"], reverse=True)
    state["operacion_finalizada"] = not progreso and bool(totals_by_opl)


def calcularProgresoOPL(totalJuegosParam=None):
    state = _load_state()
    _calcular_progreso_opl(state, totalJuegosParam)
    _save_state(state)
    return getProgresoOPL()


def getProgresoOPL():
    state = _load_state()
    return {
        "success": True,
        "progreso": state.get("opl_progreso", []),
        "operacionFinalizada": state.get("operacion_finalizada", False),
        "fecha": _now(),
    }


def getOplConfig():
    state = _load_state()
    return {"success": True, "config": state.get("opl_config", []), "opls": sorted({_as_str(r[1]) for r in state.get("opl_config", []) if len(r) > 1 and _as_str(r[1])} | {OPL_DEFAULT})}


def upsertOpl(propietario, opl):
    state = _load_state()
    prop_key = _as_str(propietario).upper()
    if not prop_key or not _as_str(opl):
        return {"success": False, "message": "Propietario y OPL son obligatorios"}
    for row in state.get("opl_config", []):
        if _as_str(row[0]).upper() == prop_key:
            row[1] = _as_str(opl)
            _save_state(state)
            return {"success": True}
    state.setdefault("opl_config", []).append([prop_key, _as_str(opl), 0])
    _save_state(state)
    return {"success": True}


def eliminarOpl(rowIdx):
    state = _load_state()
    idx = int(rowIdx) - 2
    if 0 <= idx < len(state.get("opl_config", [])):
        state["opl_config"].pop(idx)
        _save_state(state)
    return {"success": True}


def getOplPorPropietario():
    state = _load_state()
    resultado = []
    for row in state.get("opl_config", []):
        if len(row) >= 3 and int(_num(row[2])) > 0:
            resultado.append({"propietario": row[0], "opl": row[1], "juegos": int(_num(row[2]))})
    return {"success": True, "resultado": resultado, "opls": getOplConfig()["opls"]}


def getInformeDatos():
    return {"success": True, **_load_state().get("informe", {})}


def guardarInformeDatos(payload):
    state = _load_state()
    data = json.loads(payload) if isinstance(payload, str) else (payload or {})
    state["informe"] = data
    _save_state(state)
    return {"success": True}


def limpiarInformeDatos():
    state = _load_state()
    state["informe"] = _blank_state()["informe"]
    _save_state(state)
    return {"success": True}


def _total_juegos_informe(cavas):
    def inv(name):
        return sum(_num(c.get("inventario")) for c in cavas if c.get("grupo") == name)
    partes = [
        inv("V. Rojas & Blancas (V.Rojas)"),
        inv("V. Rojas & Blancas (V.Blancas)") + inv("V. Acondicionamiento"),
        inv("Patas & Manos"),
        inv("Cabezas"),
    ]
    if not partes:
        return 0
    minimo = min(partes)
    return round(minimo + sum(max(0, p - minimo) * 0.25 for p in partes), 2)


def generarInformeHTML(opciones=None):
    data = getInformeDatos()
    opts = json.loads(opciones) if isinstance(opciones, str) else (opciones or {})
    cavas = data.get("cavas", [])
    percheros = data.get("percheros", [])
    novedades = data.get("novedades", [])
    total = _total_juegos_informe(cavas)
    rows_cavas = "".join(f"<tr><td>{c.get('grupo','')}</td><td>{c.get('carros',0)}</td><td>{c.get('capPorCarro',0)}</td><td>{c.get('inventario',0)}</td></tr>" for c in cavas)
    rows_perch = "".join(f"<tr><td>{p.get('cava','')}</td><td>{p.get('blancas',0)}</td><td>{p.get('rojas',0)}</td><td>{p.get('patasManos',0)}</td><td>{p.get('cabezas',0)}</td><td>{p.get('crudas',0)}</td></tr>" for p in percheros)
    rows_nov = "".join(f"<li><strong>{n.get('cod','')}</strong> {n.get('desc','')}</li>" for n in novedades)
    body = f"""
    <h1>Informe Laboral Diario - Gestor de Visceras</h1>
    <p><strong>Fecha:</strong> {data.get('fecha','')}</p>
    <div class="kpis"><div>Total juegos: <strong>{total}</strong></div><div>Completos: <strong>{data.get('completos',0)}</strong></div><div>Incompletos: <strong>{data.get('incompletos',0)}</strong></div></div>
    """
    if opts.get("cavas", True):
        body += f"<h2>Cavas</h2><table><thead><tr><th>Grupo</th><th>Carros</th><th>Cap/Carro</th><th>Inventario</th></tr></thead><tbody>{rows_cavas}</tbody></table>"
    if opts.get("percheros", True):
        body += f"<h2>Percheros</h2><table><thead><tr><th>Cava</th><th>Blancas</th><th>Rojas</th><th>P/M</th><th>Cabezas</th><th>Crudas</th></tr></thead><tbody>{rows_perch}</tbody></table>"
    body += f"<h2>Novedades</h2><ul>{rows_nov or '<li>Sin novedades</li>'}</ul>"
    html = f"""<!doctype html><html><head><meta charset="utf-8"><title>Informe Laboral</title>
    <style>body{{font-family:Arial,sans-serif;padding:28px;color:#1f2937}}h1,h2{{color:#259c39}}table{{border-collapse:collapse;width:100%;margin:12px 0}}th{{background:#259c39;color:white}}td,th{{border:1px solid #ddd;padding:8px}}.kpis{{display:flex;gap:12px;flex-wrap:wrap}}.kpis div{{background:#f0faf3;border:1px solid #b7e4c7;border-radius:8px;padding:10px}}</style>
    </head><body>{body}</body></html>"""
    return {"success": True, "html": html}


def generarPDFDecomisos():
    resumen = getResumenDecomisos().get("resultados", [])
    rows = "".join(f"<tr><td>{r.get('id','')}</td><td>{r.get('destino','')}</td><td>{r.get('producto','')}</td></tr>" for r in resumen)
    html = f"""<!doctype html><html><head><meta charset="utf-8"><title>Decomisos</title></head><body>
    <h1>Listado de Decomisos Visceras con salida</h1><table border="1" cellpadding="6" cellspacing="0">
    <thead><tr><th>Codigo</th><th>Destino</th><th>Decomiso</th></tr></thead><tbody>{rows}</tbody></table></body></html>"""
    return {"success": True, "nombre": "Listado_Decomisos.html", "url": "data:text/html;charset=utf-8," + html}


def getHistorialPDF():
    return {"success": True, "historial": _load_state().get("historial_pdf", [])}


def cerrarOperacion():
    state = _load_state()
    progreso = state.get("opl_progreso", [])
    insertados = len(progreso)
    for row in progreso:
        item = dict(row)
        item["fecha"] = _now()
        state.setdefault("historico", []).append(item)
    state["resumen_despachos"] = {"success": True, "hayDatos": False}
    state["opl_progreso"] = []
    _save_state(state)
    return {"success": True, "insertados": insertados}


def getAniosDisponibles():
    anios = {datetime.now().year}
    for row in _load_state().get("historico", []):
        match = re.search(r"(\d{4})", _as_str(row.get("fecha")))
        if match:
            anios.add(int(match.group(1)))
    return {"success": True, "anios": sorted(anios)}


def getListaOPLsHistorico():
    state = _load_state()
    opls = sorted({_as_str(r.get("opl")) for r in state.get("historico", []) if _as_str(r.get("opl"))} | {"Todos los OPLs"})
    return {"success": True, "opls": opls}


def getKPIs(opl="Todos los OPLs", filtro=None):
    historico = _load_state().get("historico", [])
    if opl and opl != "Todos los OPLs":
        historico = [r for r in historico if r.get("opl") == opl]
    total_desp = sum(int(_num(r.get("despachados"))) for r in historico)
    total_pend = sum(int(_num(r.get("pendientes"))) for r in historico)
    kpis = {
        "vs": {"particulares": total_desp, "transcarnes": 0},
        "semana": {"despachados": total_desp, "promDiario": 0},
        "mes": {"despachados": total_desp, "tendencia": "estable"},
        "anio": {"despachados": total_desp, "operaciones": len(historico)},
        "productividad": {"promedioPorOp": total_desp, "totalOps": len(historico), "pctCompletadas": 0},
    }
    graficos = {
        "evolucion": [{"label": "Hoy", "despachados": total_desp, "pendientes": total_pend}],
        "ranking": [{"opl": r.get("opl", ""), "despachados": int(_num(r.get("despachados")))} for r in historico[:8]],
        "porDiaSemana": [{"dia": d, "actual": 0, "anterior": 0, "variacion": None, "pctActual": 0, "esHoy": False} for d in ["L", "M", "M", "J", "V", "S", "D"]],
        "eficiencia": {"completas": 0, "conPendientes": len(historico), "pctCompletas": 0, "pctPendientes": 100 if historico else 0},
    }
    backlog = [{"opl": r.get("opl", ""), "opsPendientes": 1, "totalPendientes": r.get("pendientes", 0), "eficiencia": r.get("progreso", 0), "estado": "ok"} for r in historico if _num(r.get("pendientes")) > 0]
    return {"success": True, "kpis": kpis, "graficos": graficos, "backlog": backlog}


def getListaOPLsParaPlanilla():
    state = _load_state()
    return sorted({_as_str(row[1]) for row in state.get("opl_config", []) if len(row) > 1 and _as_str(row[1])} | {OPL_DEFAULT})


def consolidarDatos():
    return {"success": True, "procesados": contarJuegosViscerales().get("total", 0), "turno": _detectar_turno_desde_datos(_load_state())}


def generarPlanillaPuntos(opl):
    return {"success": True, "opl": opl, "totalOPL": 0, "porcentaje": 0, "zonas": []}


def generarHTMLPlanillaPDF(opl):
    html = f"<!doctype html><html><body><h1>Planilla de puntos</h1><p>OPL: {opl}</p></body></html>"
    return {"success": True, "html": html}


def mostrarResumenGeneral():
    return {"success": True, "resumen": [], "totalGeneral": 0}


def generarReporteOPL(opl, filtro=None):
    html = f"<!doctype html><html><body><h1>Reporte OPL</h1><p>{opl}</p></body></html>"
    return {"success": True, "html": html}


def dispatch(function_name, args):
    allowed = globals().get(function_name)
    if not callable(allowed) or function_name.startswith("_"):
        return {"success": False, "message": f"Funcion no implementada: {function_name}"}
    return allowed(*(args or []))
